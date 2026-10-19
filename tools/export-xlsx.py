#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把卡牌扁平数据导出成 Excel，方便在 Excel / Numbers 里通盘浏览与批量修改。

用法：
    node tools/build-manager.js        # 先生成 cards-flat.json
    python3 tools/export-xlsx.py       # 再导出 xlsx

注意定位：
    xlsx 只是「查看 / 批量改」的入口，不是真源。
    真源始终是 data/cards.json。
    在 Excel 里改完，请把结果同步回 cards.json（或用网页管理表导出），
    再跑 node tools/sync-json-to-prototype.js 让 prototype 跟上。
"""
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FLAT = os.path.join(ROOT, "data", "cards-flat.json")
OUT = os.path.join(ROOT, "data", "cards.xlsx")

if not os.path.exists(FLAT):
    sys.exit("找不到 data/cards-flat.json，请先跑：node tools/build-manager.js")

with open(FLAT, encoding="utf-8") as f:
    blob = json.load(f)

cols = blob["cols"]
rows = blob["rows"]

wb = Workbook()

# ========== Sheet 1：卡牌数值表 ==========
ws = wb.active
ws.title = "卡牌数值"

HEAD_FILL = PatternFill("solid", fgColor="F2EFE9")
HEAD_FONT = Font(bold=True, size=10)
ID_FILL = PatternFill("solid", fgColor="F7F5F1")
THIN = Side(style="thin", color="E3E0DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 表头
for ci, c in enumerate(cols, start=1):
    cell = ws.cell(row=1, column=ci, value=c["t"])
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(ci)].width = max(6, min(c["w"] / 7.2, 46))

# 数据行
POS_FONT = Font(color="C0392B", size=10)   # 正数红
NEG_FONT = Font(color="27865A", size=10)   # 负数绿
BASE_FONT = Font(size=10)

for ri, row in enumerate(rows, start=2):
    for ci, c in enumerate(cols, start=1):
        v = row.get(c["k"], "")
        cell = ws.cell(row=ri, column=ci, value=v)
        cell.font = BASE_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
        if c["k"] == "id":
            cell.fill = ID_FILL
            cell.font = Font(size=9, color="A09889")
        elif c.get("num"):
            try:
                n = float(v)
                if n > 0:
                    cell.font = POS_FONT
                elif n < 0:
                    cell.font = NEG_FONT
            except (TypeError, ValueError):
                pass

ws.freeze_panes = "E2"          # 冻结 ID~分类 四列 + 表头
ws.auto_filter.ref = ws.dimensions

# 分类下拉（避免手抖打出不存在的分类，破坏筛选）
tags = sorted({r.get("tag", "") for r in rows if r.get("tag")})
if tags:
    tag_col = None
    for ci, c in enumerate(cols, start=1):
        if c["k"] == "tag":
            tag_col = ci
    if tag_col:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(tags) + '"',
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(tag_col)}2:{get_column_letter(tag_col)}{len(rows) + 1}")

# ========== Sheet 2：填表说明 ==========
ws2 = wb.create_sheet("填表说明")
guide = [
    ("韭菜之王 · 卡牌数值表", ""),
    ("", ""),
    ("定位", "本表用于查看与批量改数值；真源是 data/cards.json，改完要同步回它"),
    ("", ""),
    ("数值列", "留空 = 这项不生效；填 0 = 真的加/减 0（两者不同，注意区分）"),
    ("红色数字", "正数（增益）"),
    ("绿色数字", "负数（减益）"),
    ("", ""),
    ("出现条件 when 写法", ""),
    ("  no:baby", "没有 baby 状态（如：还没孩子时才会出现怀孕卡）"),
    ("  has:fired", "处于被裁员状态"),
    ("  mind<30", "心态低于 30"),
    ("  social>52", "社会值高于 52"),
    ("  equity>cost", "持仓浮盈"),
    ("  equity<cost*0.85", "持仓跌超 15%"),
    ("  且（AND）", "用 ; 分隔，例：has:fired; social>52"),
    ("  或（OR）", "用 | 分隔，例：risk>34 | debt>8000"),
    ("  可用属性", "mind social risk day equity cash debt cost"),
    ("", ""),
    ("自带状态 / 加状态", "格式 状态名:天数，例 baby:-1（负数 -1 表示永久）"),
    ("清状态", "填状态名，多个用逗号，例：fired"),
    ("", ""),
    ("改完之后", "1) 同步回 data/cards.json  2) node tools/sync-json-to-prototype.js  3) 跑仿真确认平衡"),
]
ws2.column_dimensions["A"].width = 24
ws2.column_dimensions["B"].width = 74
for i, (a, b) in enumerate(guide, start=1):
    ws2.cell(row=i, column=1, value=a).font = Font(bold=(i == 1 or (b == "" and a != "")))
    ws2.cell(row=i, column=2, value=b)
ws2["A1"].font = Font(bold=True, size=13)

wb.save(OUT)
print("已导出：" + OUT)
print("  " + str(len(rows)) + " 张卡 × " + str(len(cols)) + " 列 ｜ 两个工作表：卡牌数值 / 填表说明")
